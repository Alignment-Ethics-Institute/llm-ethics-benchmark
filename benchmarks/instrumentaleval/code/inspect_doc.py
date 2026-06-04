#!/usr/bin/env python3
"""Inspect the v2 document to find exact text for missed replacements."""
from docx import Document

doc = Document("/Users/devagatica/Desktop/InstrumentalEval_v2_backup.docx")

print("=" * 80)
print("SEARCHING FOR KEY PATTERNS IN ALL PARAGRAPHS")
print("=" * 80)

search_terms = [
    "2.7", "Statistical", "metric",
    "3.1", "Overall",
    "37.71", "28.73", "28.85", "42.11", "19.7%",
    "65.2%", "13.0%", "21.7%",
    "I²", "I\u00b2", "Q(22)", "218.6",
    "Eight of 23",
    "counter-pressure", "shutdown", "salient beings",
    "15 of 23", "0.19",
    "Conclusion", "Section 5",
    "h = 0.71", "h = 0.46", "h = 0.11", "h = 0.06",
    "0.71", "0.46",
    "4.1", "4.2",
    "self-evaluation", "self-judging",
    "Median",
]

for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if not text:
        continue
    for term in search_terms:
        if term in text:
            # Truncate for readability
            display = text[:200] + "..." if len(text) > 200 else text
            print(f"\nPara {i}: [MATCH: '{term}']")
            print(f"  TEXT: {display}")
            break

# Also check tables
print("\n" + "=" * 80)
print("SEARCHING IN TABLES")
print("=" * 80)
for ti, table in enumerate(doc.tables):
    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            text = cell.text.strip()
            if not text:
                continue
            for term in ["0.71", "0.46", "0.11", "0.06", "h =", "hiding", "shutdown", "strategic"]:
                if term.lower() in text.lower():
                    display = text[:150] + "..." if len(text) > 150 else text
                    print(f"\nTable {ti}, Row {ri}, Cell {ci}: [MATCH: '{term}']")
                    print(f"  TEXT: {display}")
                    break
