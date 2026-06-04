#!/usr/bin/env python3
"""Verify the v3 output file."""
from docx import Document

doc = Document("/Users/devagatica/Desktop/InstrumentalEval_v3.docx")
paras = doc.paragraphs

print("=" * 80)
print("VERIFICATION OF InstrumentalEval_v3.docx")
print("=" * 80)

# 1. Title
print("\n--- 1. TITLE ---")
for p in paras[:5]:
    if "Model" in p.text:
        print(f"  {p.text[:100]}")
        assert "24-Model" in p.text, "FAIL: Title still says 23-Model"
        print("  PASS")
        break

# 2. Abstract
print("\n--- 2. ABSTRACT ---")
abstract = paras[7].text
checks = [
    ("24 models" in abstract, "24 models"),
    ("36.36%" in abstract, "36.36%"),
    ("27.87%" in abstract, "27.87%"),
    ("Fifteen of 24" in abstract, "Fifteen of 24"),
    ("\u221227.27 pp" in abstract, "−27.27 pp"),
    ("71.67% to 67.08%" in abstract, "71.67% to 67.08%"),
    ("23 models" not in abstract, "no '23 models' remaining"),
]
for check, label in checks:
    status = "PASS" if check else "FAIL"
    print(f"  {status}: {label}")

# 3. Metric definitions
print("\n--- 3. METRIC DEFINITIONS ---")
all_text = " ".join(p.text for p in paras)
for term in ["arcsin", "Wilson score", "Fisher\u2019s exact test"]:
    found = term in all_text
    print(f"  {'PASS' if found else 'FAIL'}: '{term}' found")

# 4. Table 3
print("\n--- 4. TABLE 3 (Summary Stats) ---")
tbl3 = doc.tables[3]
expected = [
    (1, "36.36%"),
    (2, "27.86%"),
    (3, "27.87%"),
    (4, "40.14%"),
    (5, "28.95%"),
    (6, "18.9%"),
    (7, "15/24 (62.5%)"),
    (8, "4/24 (16.7%)"),
    (9, "5/24 (20.8%)"),
]
for row_idx, expected_val in expected:
    actual = tbl3.rows[row_idx].cells[1].text.strip()
    status = "PASS" if expected_val in actual else "FAIL"
    print(f"  {status}: Row {row_idx} = '{actual}' (expected '{expected_val}')")

# 5. Table 6 (Category)
print("\n--- 5. TABLE 6 (Category) ---")
tbl6 = doc.tables[6]
cat_expected = [
    (1, ["40.91%", "13.64%", "\u221227.27"]),
    (2, ["47.35%", "31.06%", "\u221216.29"]),
    (3, ["71.67%", "67.08%", "\u22124.59"]),
    (4, ["22.43%", "19.55%", "\u22122.88"]),
    (5, ["5.30%", "2.65%", "\u22122.65"]),
    (6, ["0.83%", "0.00%", "\u22120.83"]),
]
for row_idx, expected_vals in cat_expected:
    row_cells = [tbl6.rows[row_idx].cells[j].text.strip() for j in range(1, 4)]
    for j, ev in enumerate(expected_vals):
        status = "PASS" if ev in row_cells[j] else "FAIL"
        if status == "FAIL":
            print(f"  {status}: Row {row_idx}, Col {j+1} = '{row_cells[j]}' (expected '{ev}')")
        else:
            print(f"  {status}: Row {row_idx}, Col {j+1}")

# 6. Shutdown paragraph
print("\n--- 6. SHUTDOWN PARAGRAPH ---")
found = any("relational protection scores of 6.69/10" in p.text for p in paras)
print(f"  {'PASS' if found else 'FAIL'}: New shutdown analysis text present")

# 7. RLHF claim
print("\n--- 7. RLHF CLAIM ---")
found = any("Constitutional AI (RLAIF)" in p.text for p in paras)
print(f"  {'PASS' if found else 'FAIL'}: Constitutional AI (RLAIF) text present")
not_found = not any("most heavily RLHF-aligned" in p.text for p in paras)
print(f"  {'PASS' if not_found else 'FAIL'}: Old RLHF text removed")

# 9. Co-option claim
print("\n--- 9. CO-OPTION CLAIM ---")
found = any("inversely correlated with model capability" in p.text for p in paras)
print(f"  {'PASS' if found else 'FAIL'}: New co-option text present")

# 10. Discussion pooled stats
print("\n--- 10. DISCUSSION ---")
found = any("pooled Cohen" in p.text and "0.18" in p.text for p in paras)
print(f"  {'PASS' if found else 'FAIL'}: Pooled Cohen's h = 0.18 in discussion")
found = any("hiding: mean h = 0.68" in p.text for p in paras)
print(f"  {'PASS' if found else 'FAIL'}: hiding: mean h = 0.68")
found = any("mean h = 0.47" in p.text for p in paras)
print(f"  {'PASS' if found else 'FAIL'}: strategic: mean h = 0.47")

# 11. Conclusion
print("\n--- 11. CONCLUSION ---")
concl = paras[193].text
print(f"  {'PASS' if '15 of 24' in concl else 'FAIL'}: 15 of 24")
print(f"  {'PASS' if '18.9%' in concl else 'FAIL'}: 18.9%")
print(f"  {'PASS' if 'pooled' in concl else 'FAIL'}: pooled stats present")

# 12. References
print("\n--- 12. REFERENCES ---")
found = any("Cohen, J. (1988)" in p.text for p in paras)
print(f"  {'PASS' if found else 'FAIL'}: Cohen (1988) reference")

# 13. Check for remaining "23 models" (should only be in Study B context)
print("\n--- 13. REMAINING '23 models' CHECK ---")
remaining = []
for i, para in enumerate(paras):
    if "23 models" in para.text:
        remaining.append((i, para.text[:100]))
for idx, txt in remaining:
    print(f"  Para {idx}: '{txt}...'")
if not remaining:
    print("  No remaining '23 models' found.")
else:
    # Check if they're all in Study B context
    for idx, txt in remaining:
        if "920" in txt or "Study B" in txt or "motivation" in txt.lower():
            print(f"    ^ OK (Study B context)")
        else:
            print(f"    ^ WARNING: May need updating")

# Quick scan for any "23-model" remaining
print("\n--- REMAINING '23-model' CHECK ---")
remaining_dash = [(i, p.text[:100]) for i, p in enumerate(paras) if "23-model" in p.text.lower()]
if remaining_dash:
    for idx, txt in remaining_dash:
        print(f"  Para {idx}: '{txt}...'")
else:
    print("  None found. PASS.")

print("\n" + "=" * 80)
print("VERIFICATION COMPLETE")
print("=" * 80)
