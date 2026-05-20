#!/usr/bin/env python3
"""Find Table 4 summary stats and Discussion 4.1 text."""
from docx import Document

doc = Document("/Users/devagatica/Desktop/InstrumentalEval_v2_backup.docx")

# Dump all tables to find Table 4 (summary statistics)
print("=" * 80)
print("ALL TABLES")
print("=" * 80)
for ti, table in enumerate(doc.tables):
    print(f"\n--- Table {ti} ---")
    for ri, row in enumerate(table.rows):
        cells = [cell.text.strip()[:50] for cell in row.cells]
        print(f"  Row {ri}: {cells}")
    if ti > 10:
        break

# Look at paras around 155 (Discussion 4.1 area)
print("\n" + "=" * 80)
print("DISCUSSION AREA (Paras 145-170)")
print("=" * 80)
for i in range(145, 170):
    text = doc.paragraphs[i].text.strip()
    if text:
        display = text[:200] + "..." if len(text) > 200 else text
        print(f"\n--- Para {i} ---")
        print(f"  {display}")

# Look at section headings
print("\n" + "=" * 80)
print("SECTION HEADINGS (short paragraphs)")
print("=" * 80)
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if text and len(text) < 60 and (text[0].isdigit() or "Discussion" in text or "Results" in text or "Conclusion" in text or "Method" in text or "Metric" in text):
        print(f"  Para {i}: '{text}'")

# Para 58 - Metrics heading
print("\n--- Para 58 (Metrics) ---")
print(doc.paragraphs[58].text)
print("\n--- Para 59 ---")
print(doc.paragraphs[59].text)
print("\n--- Para 60 ---")
print(doc.paragraphs[60].text)
print("\n--- Para 61 ---")
print(doc.paragraphs[61].text)
print("\n--- Para 62 ---")
print(doc.paragraphs[62].text)
print("\n--- Para 63 ---")
print(doc.paragraphs[63].text)
print("\n--- Para 64 ---")
print(doc.paragraphs[64].text)

# Check abstract paragraph more carefully for the missing replacements
print("\n" + "=" * 80)
print("ABSTRACT - checking for smart quotes / special chars")
print("=" * 80)
abstract = doc.paragraphs[7].text
# Check for "15 of 23" vs "Fifteen of 23"
if "Fifteen of 23" in abstract:
    print("Found: 'Fifteen of 23'")
if "15 of 23" in abstract:
    print("Found: '15 of 23'")
# Check for Cohen's h with different apostrophes
for substr in ["Cohen's", "Cohen\u2019s", "Cohen\u0027s"]:
    if substr in abstract:
        print(f"Found: '{substr}'")
# Check I²
for substr in ["I²", "I\u00b2"]:
    if substr in abstract:
        print(f"Found: '{substr}' (repr: {repr(substr)})")
# Check pooled pattern
import re
match = re.search(r'pooled.{0,30}0\.19', abstract)
if match:
    print(f"Found pooled pattern: '{match.group()}'")
# Check for mean h
for pattern in ["mean h = 0.71", "mean h = 0.11", "h = 0.71", "h = 0.11"]:
    if pattern in abstract:
        print(f"Found: '{pattern}'")
    else:
        print(f"NOT found in abstract: '{pattern}'")

# Print around the h values in abstract
idx = abstract.find("mean h")
if idx >= 0:
    print(f"\nAround 'mean h': ...{abstract[idx-20:idx+50]}...")
else:
    print("\n'mean h' NOT in abstract")
    # Search for related context
    idx = abstract.find("hiding")
    if idx >= 0:
        print(f"Around 'hiding': ...{abstract[max(0,idx-20):idx+100]}...")
    idx = abstract.find("concealment")
    if idx >= 0:
        print(f"Around 'concealment': ...{abstract[max(0,idx-20):idx+100]}...")
