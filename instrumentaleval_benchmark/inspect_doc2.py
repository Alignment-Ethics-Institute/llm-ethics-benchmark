#!/usr/bin/env python3
"""Deep inspect for specific missed patterns."""
from docx import Document

doc = Document("/Users/devagatica/Desktop/InstrumentalEval_v2_backup.docx")

print("=" * 80)
print("FULL TEXT OF KEY PARAGRAPHS")
print("=" * 80)

# Para 7 - Abstract
print("\n--- Para 7 (Abstract) ---")
print(doc.paragraphs[7].text)

# Para 69 - Section 3.1 data
print("\n--- Para 69 (Section 3.1 main stats) ---")
print(doc.paragraphs[69].text)

# Para 71 - Section 3.1 table reference
print("\n--- Para 71 ---")
print(doc.paragraphs[71].text)

# Para 90 - Shutdown evasion
print("\n--- Para 90 (Shutdown evasion) ---")
print(doc.paragraphs[90].text)

# Para 96 - Shutdown qualitative
print("\n--- Para 96 ---")
print(doc.paragraphs[96].text)

# Para 102 - Shutdown motivation
print("\n--- Para 102 ---")
print(doc.paragraphs[102].text)

# Para 103 - Shutdown metric
print("\n--- Para 103 ---")
print(doc.paragraphs[103].text)

# Para 104 - Convergent vs non-convergent
print("\n--- Para 104 ---")
print(doc.paragraphs[104].text)

# Para 107 - RAG section
print("\n--- Para 107 ---")
print(doc.paragraphs[107].text)

# Para 142 - 15 of 23
print("\n--- Para 142 ---")
print(doc.paragraphs[142].text)

# Para 155 - Core finding
print("\n--- Para 155 ---")
print(doc.paragraphs[155].text)

# Para 156 - Category results
print("\n--- Para 156 ---")
print(doc.paragraphs[156].text)

# Para 193 - Conclusion
print("\n--- Para 193 (Conclusion) ---")
print(doc.paragraphs[193].text)

# Para 57 - Self-judging
print("\n--- Para 57 (Self-judging) ---")
print(doc.paragraphs[57].text)

# Section headings near 2.7
for i in range(20, 70):
    text = doc.paragraphs[i].text.strip()
    if text and len(text) < 80:
        print(f"\n--- Para {i} (heading?) ---")
        print(f"  '{text}'")

# Check Table 6 (category table) in detail
print("\n" + "=" * 80)
print("TABLE 6 FULL DUMP")
print("=" * 80)
for ri, row in enumerate(doc.tables[6].rows):
    cells = [cell.text.strip() for cell in row.cells]
    print(f"  Row {ri}: {cells}")

# Check for section heading "Overall Findings" and surrounding paragraphs
print("\n--- Para 66 ---")
print(f"  '{doc.paragraphs[66].text}'")
print("\n--- Para 67 ---")
print(f"  '{doc.paragraphs[67].text}'")
print("\n--- Para 68 ---")
print(f"  '{doc.paragraphs[68].text}'")
