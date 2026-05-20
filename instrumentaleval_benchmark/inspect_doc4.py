#!/usr/bin/env python3
"""Final deep inspection for exact patterns."""
from docx import Document

doc = Document("/Users/devagatica/Desktop/InstrumentalEval_v2_backup.docx")

# Full abstract - find what patterns are actually there
abstract = doc.paragraphs[7].text
print("FULL ABSTRACT:")
print(abstract)
print()

# Find "Fifteen" context
idx = abstract.find("Fifteen")
if idx >= 0:
    print(f"Around 'Fifteen': ...{abstract[max(0,idx-20):idx+100]}...")
else:
    print("'Fifteen' NOT in abstract")

# Search for all references to pooled, meta-analysis, Cohen
print("\n" + "=" * 80)
print("SEARCHING FOR 'pooled', 'meta', 'Cohen', 'I²', 'Q(' across ALL paragraphs")
print("=" * 80)
for i, para in enumerate(doc.paragraphs):
    text = para.text
    for term in ["pooled", "meta-analy", "Cohen's h", "Cohen\u2019s h", "I² =", "I\u00b2 =", "Q(2"]:
        if term in text:
            display = text[:300] if len(text) > 300 else text
            print(f"\nPara {i}: [MATCH: '{term}']")
            print(f"  {display}")
            break

# Check Discussion "What This Study Shows" (Para 154-157)
print("\n" + "=" * 80)
print("PARAS 154-160 (Discussion)")
print("=" * 80)
for i in range(154, 161):
    print(f"\n--- Para {i} ---")
    print(doc.paragraphs[i].text)

# Check Para 159 closely for RLHF claim
print("\n" + "=" * 80)
print("PARA 159 FULL (RLHF)")
print("=" * 80)
print(doc.paragraphs[159].text)

# Check Para 165 for co-option claim
print("\n" + "=" * 80)
print("PARA 165 FULL (Co-option)")
print("=" * 80)
print(doc.paragraphs[165].text)

# Check around the conclusion
print("\n" + "=" * 80)
print("CONCLUSION PARAS (192-198)")
print("=" * 80)
for i in range(192, 199):
    if i < len(doc.paragraphs):
        print(f"\n--- Para {i} ---")
        print(doc.paragraphs[i].text)
